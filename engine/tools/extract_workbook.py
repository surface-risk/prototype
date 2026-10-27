"""Extract methodology config, fictional sample inputs and expected values from the v0.5 workbook.

Usage: python tools/extract_workbook.py path/to/Surface-Risk_Prototype_Framework_v0.5.xlsx

Needs openpyxl (only this tool does; the engine itself is stdlib-only).
Writes:
  config/methodology_v0.5.json
  sample_data/{firm_profile,staff_responses,mgmt_responses,check_results,peers}.json
  tests/expected_sample.json   (values the workbook calculated, used by the tests)
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]


def rows(ws, values=True):
    for row in ws.iter_rows(values_only=values):
        yield list(row)


def num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


def main(path: str) -> None:
    wb = openpyxl.load_workbook(path, data_only=True)

    # ---------- parameters ----------
    sm = wb["Scoring_Model"]
    p = {r[0]: r[1] for r in rows(sm) if r[0] and isinstance(r[1], (int, float))}
    params = {
        "min_reporting_group": p["Minimum reporting group"],
        "blind_spot_threshold": p["Blind-spot threshold (points)"],
        "speeder_threshold_x_median": p["Speeder threshold (x median time)"],
        "archetype_low": p["Archetype boundary - low"],
        "archetype_high": p["Archetype boundary - high"],
        "max_non_survey_weight": p["Maximum non-survey weight per area"],
        "tier_points": {"H": p["Importance points - High"], "M": p["Importance points - Medium"], "L": p["Importance points - Low"]},
        "shrinkage_k": p["Shrinkage strength (pseudo-respondents)"],
        "z95": p["Confidence multiplier"],
        "min_data_points": p["Minimum data points per area"],
        "exposure_baseline": p["Exposure baseline (0-100 risk)"],
        "penalty_per_point": p["Penalty per point of exposure above baseline"],
        "recoverable_share": p["Recoverable share of hidden capacity"],
    }
    bands = [{"label": sm.cell(r, 1).value, "from": sm.cell(r, 2).value, "meaning": sm.cell(r, 3).value} for r in range(32, 37)]
    archetypes = [{"key": sm.cell(r, 1).value, "name": sm.cell(r, 2).value, "description": sm.cell(r, 3).value} for r in range(40, 49)]

    def opts(r0, r1):
        return [{"answer": str(sm.cell(r, 1).value), "points": sm.cell(r, 2).value} for r in range(r0, r1 + 1)]

    mgmt_scoring = {
        "MC2": opts(52, 55), "MC3": opts(57, 61), "MC5": opts(63, 67), "MC6": opts(69, 73), "MC7": opts(75, 78),
        "MC4": "percent of material providers attributed to the owner recorded at set-up (capped at 100)",
        "MC8": "50 if contact matches the registered first responder + 50 if stored offline",
    }
    hierarchy = [{"level_c": sm.cell(r, 1).value, "level_b": sm.cell(r, 2).value, "level_a": sm.cell(r, 3).value} for r in range(83, 91)]
    default_level = [
        {"max_headcount_exclusive": 30, "level": "A"},
        {"max_headcount_exclusive": 151, "level": "B"},
        {"max_headcount_exclusive": None, "level": "C"},
    ]

    # ---------- areas, tiers, overrides ----------
    ru = wb["Reporting_Units"]
    bl_cols = ["BL01", "BL02", "BL03", "BL04", "BL05", "BL06", "BL07", "BL08", "BL09", "INS-L", "INS-N"]
    areas = []
    for r in range(5, 25):
        areas.append({
            "code": ru.cell(r, 1).value, "name": ru.cell(r, 2).value, "orx_l1": ru.cell(r, 3).value,
            "orx_l2": ru.cell(r, 4).value, "loss_tier": ru.cell(r, 5).value, "perf_tier": ru.cell(r, 6).value,
            "rationale": ru.cell(r, 7).value, "loss_overrides": {}, "perf_overrides": {},
        })
    by_code = {a["code"]: a for a in areas}

    def read_overrides(r0, key):
        for r in range(r0, r0 + 20):
            code = ru.cell(r, 1).value
            for i, bl in enumerate(bl_cols):
                v = ru.cell(r, 7 + i).value
                if v:
                    by_code[code][key][bl] = v
    # find section rows
    sec = {}
    for r in range(1, ru.max_row + 1):
        v = ru.cell(r, 1).value
        if isinstance(v, str) and v.startswith("B. Loss-tier"):
            sec["loss"] = r + 2
        if isinstance(v, str) and v.startswith("C. Performance-tier"):
            sec["perf"] = r + 2
    read_overrides(sec["loss"], "loss_overrides")
    read_overrides(sec["perf"], "perf_overrides")

    business_lines = []
    for r in rows(wb["Business_Lines"]):
        if r[0] and (str(r[0]).startswith("BL") or str(r[0]).startswith("INS")):
            business_lines.append({"code": r[0], "label": r[1], "orx_l2": r[2], "basel": r[3], "guidance": r[4]})

    # ---------- statements ----------
    qb = wb["Question_Bank"]
    statements = []
    for r in rows(qb):
        if r[0] and str(r[0]).startswith("Q") and r[2] and r[4]:
            statements.append({"id": r[0], "cause_code": r[1], "area": r[2], "text": r[4], "format": r[5], "reverse": r[6] == "Y"})
    # order statements as in Scored (engine order does not matter, but keep it stable)
    sc = wb["Scored"]
    order = [sc.cell(3, c).value for c in range(5, 65)]
    statements.sort(key=lambda s: order.index(s["id"]))

    other_q = {}
    section = None
    for r in rows(qb):
        a = r[0]
        if isinstance(a, str) and len(a) > 2 and a[1] == "." and a[0] in "CDEFGHI":
            section = a
            continue
        if section and a and not str(a).startswith("ID"):
            other_q.setdefault(section, []).append({"id": a, "label": r[3], "text": r[4], "options": r[5], "scoring": r[8]})

    # ---------- checks and exposures ----------
    ce = wb["Checks_Exposure"]
    checks = []
    for r in range(5, 16):
        row = [ce.cell(r, c).value for c in range(1, 13)]
        checks.append({"id": row[0], "integration": row[1], "name": row[2], "area": row[3], "weight": row[4], "metric": row[5],
                       "direction": row[6], "thresholds": [num(x) for x in row[7:11]], "source": row[11]})
    exposures = []
    for r in range(19, 25):
        row = [ce.cell(r, c).value for c in range(1, 11)]
        exposures.append({"id": row[0], "integration": row[1], "name": row[2], "area": row[3], "index": row[4],
                          "footprint_weights": dict(zip(["OFF", "AFF", "CLI", "PRV", "CPT"], row[5:10]))})

    cr = wb["Country_Risk"]
    idx_types = [cr.cell(4, c).value for c in range(2, 8)]
    countries = {}
    for r in range(6, 20):
        countries[cr.cell(r, 1).value] = dict(zip(idx_types, [cr.cell(r, c).value for c in range(2, 8)]))

    config = {
        "methodology_version": "0.5",
        "source": "Surface-Risk_Prototype_Framework_v0.5.xlsx",
        "notes": "Importance tiers are expert first-pass estimates pending practitioner panel review. Country risk values are ILLUSTRATIVE only.",
        "parameters": params,
        "answer_anchors": {"how_true": [0, 25, 50, 75, 100], "how_often": [0, 25, 50, 75, 100], "yes_no": {"Yes": 100, "No": 0},
                            "reverse_rule": "warning-sign items are scored 100 - answer; Don't know = excluded"},
        "bands": bands,
        "archetypes": archetypes,
        "areas": areas,
        "business_lines": business_lines,
        "statements": statements,
        "other_questions": other_q,
        "management_scoring": mgmt_scoring,
        "department_hierarchy": hierarchy,
        "default_reporting_level": default_level,
        "checks": checks,
        "exposures": exposures,
        "country_risk": {"illustrative": True, "index_types": idx_types, "countries": countries},
    }
    (ROOT / "config" / "methodology_v0.5.json").write_text(json.dumps(config, indent=2, ensure_ascii=False))

    # ---------- sample firm ----------
    sp = wb["Sample_Profile"]
    bl_mix = {sp.cell(r, 2).value: sp.cell(r, 3).value for r in range(14, 25)}
    fp_keys = ["OFF", "AFF", "CLI", "PRV", "CPT"]
    footprint = {k: {} for k in fp_keys}
    for r in range(29, 43):
        country = sp.cell(r, 1).value
        for i, k in enumerate(fp_keys):
            v = sp.cell(r, 2 + i).value
            if v:
                footprint[k][country] = v
    providers = []
    for r in range(47, 52):
        providers.append({"name": sp.cell(r, 1).value, "service": sp.cell(r, 2).value, "provider_country": sp.cell(r, 4).value,
                          "data_country": sp.cell(r, 5).value, "criticality": sp.cell(r, 6).value, "owner": sp.cell(r, 7).value})
    profile = {
        "fictional": True, "name": sp.cell(3, 2).value, "centre": sp.cell(4, 2).value, "headcount": sp.cell(5, 2).value,
        "loaded_cost_usd": sp.cell(6, 2).value, "date": sp.cell(7, 2).value, "first_contact": sp.cell(8, 2).value,
        "domains": sp.cell(9, 2).value, "lei": sp.cell(10, 2).value, "business_line_mix": bl_mix, "footprint": footprint,
        "providers": providers,
    }
    (ROOT / "sample_data" / "firm_profile.json").write_text(json.dumps(profile, indent=2, ensure_ascii=False))

    ss = wb["Sample_Staff"]
    hdr = [c.value for c in ss[3]]
    staff = []
    for r in range(4, ss.max_row + 1):
        vals = [c.value for c in ss[r]]
        if not vals[0]:
            continue
        d = dict(zip(hdr, vals))
        staff.append({
            "id": d["Resp"], "department_c": d["Department (Level C)"], "business_line": d["Business line"],
            "seniority": d["Seniority"], "minutes": d["Minutes"],
            "answers": {q["id"]: (None if d[q["id"]] in (None, "") else float(d[q["id"]])) for q in statements},
            "rework_share": d["A2"], "sole_performer_activities": d["A3"], "unclear_activities": d["A4"],
        })
    (ROOT / "sample_data" / "staff_responses.json").write_text(json.dumps(staff, indent=1, ensure_ascii=False))

    sg = wb["Sample_Mgmt"]
    hdr = [c.value for c in sg[3]]
    mgmt = []
    for r in range(4, 8):
        d = dict(zip(hdr, [c.value for c in sg[r]]))
        mgmt.append({
            "role": d["Manager"],
            "answers": {q["id"]: (None if d[q["id"]] in (None, "") else float(d[q["id"]])) for q in statements},
            "mc1_top5": [d[f"MC1 p{i}"] for i in range(1, 6)],
            "mc2": d["MC2"], "mc3": str(d["MC3"]), "mc4_correct_owners": d["MC4"], "mc5": d["MC5"], "mc6": d["MC6"], "mc7": d["MC7"],
            "mc8_contact": d["MC8 contact"], "mc8_offline": d["MC8 offline"] == "Y",
            "alignment": {f"MA{i}": d[f"MA{i}"] for i in range(1, 9)},
        })
    (ROOT / "sample_data" / "mgmt_responses.json").write_text(json.dumps(mgmt, indent=1, ensure_ascii=False))

    si = wb["Sample_Inputs"]
    check_results = {}
    for r in range(5, 14):  # C01-C09 raw results (A01/A02 are computed from the survey)
        cid = si.cell(r, 1).value
        check_results[cid] = {"run": si.cell(r, 4).value == "Y", "result": num(si.cell(r, 5).value)}
    (ROOT / "sample_data" / "check_results.json").write_text(json.dumps(check_results, indent=1))
    peers = {si.cell(r, 1).value: si.cell(r, 3).value for r in range(30, 50)}
    (ROOT / "sample_data" / "peers.json").write_text(json.dumps({"illustrative": True, "area_scores": peers}, indent=1))

    # ---------- expected values (workbook-calculated) ----------
    ca = wb["Calc"]
    exp_areas = {}
    cols = {"items": 4, "data_points": 5, "survey": 7, "sd": 8, "n": 9, "ci": 10, "check_w": 11, "check_pts": 12, "prep": 13,
            "exposure": 14, "penalty": 15, "score": 16, "band": 17, "mgmt_survey": 18, "mgmt": 19, "gap": 20, "flag": 21,
            "peer": 22, "loss_wt": 24, "perf_wt": 25, "priority": 28}
    grp_cols = {"Front office": 31, "Operations": 36, "Control functions": 41, "Support functions": 46,
                "BL09 Private Banking & Wealth": 51, "BL07 Asset Management": 56, "Firm-wide / shared": 61}
    for r in range(82, 102):
        code = ca.cell(r, 1).value
        e = {k: ca.cell(r, c).value for k, c in cols.items()}
        e["groups"] = {g: {"raw": ca.cell(r, c).value, "n": ca.cell(r, c + 1).value, "shrunk": ca.cell(r, c + 2).value,
                           "score": ca.cell(r, c + 3).value, "flag": ca.cell(r, c + 4).value} for g, c in grp_cols.items()}
        exp_areas[code] = e
    groups = {}
    for r in range(112, 122):
        groups[ca.cell(r, 1).value] = {"n": ca.cell(r, 3).value, "loss": ca.cell(r, 4).value, "perf": ca.cell(r, 5).value,
                                       "archetype": ca.cell(r, 8).value, "flags": ca.cell(r, 11).value}
    headline = {ca.cell(r, 1).value: ca.cell(r, 2).value for r in range(154, 182)}
    managers = {ca.cell(r, 1).value: [ca.cell(r, c).value for c in range(2, 11)] for r in range(137, 142)}
    pri = [[ca.cell(r, c).value for c in range(2, 12)] for r in range(127, 132)]
    exposures_v = {si.cell(r, 1).value: si.cell(r, 10).value for r in range(20, 26)}
    check_scores = {si.cell(r, 1).value: num(si.cell(r, 11).value) for r in range(5, 16)}
    check_results_all = {si.cell(r, 1).value: num(si.cell(r, 5).value) for r in range(5, 16)}
    expected = {"areas": exp_areas, "groups": groups, "headline": headline, "managers": managers, "priorities": pri,
                "exposures": exposures_v, "check_scores": check_scores, "check_results": check_results_all,
                "equal_weight_index": ca.cell(123, 4).value, "equal_weight_archetype": ca.cell(123, 8).value,
                "top3_overlap": ca.cell(132, 4).value, "robustness": ca.cell(133, 4).value}
    (ROOT / "tests" / "expected_sample.json").write_text(json.dumps(expected, indent=1, default=str))
    print("extracted", len(statements), "statements,", len(staff), "staff,", len(mgmt), "managers")


if __name__ == "__main__":
    main(sys.argv[1])
